"""
Turns an uploaded PDF, PPTX, or video transcript into plain text pages/slides.

Video is handled as a transcript string (either pasted directly, or produced
upstream by an ASR step such as Whisper — out of scope for this module, kept
as a clean seam via `parse_video_transcript`).
"""
from __future__ import annotations

from dataclasses import dataclass

import pypdf
from pptx import Presentation


@dataclass
class ParsedUnit:
    """One page (PDF) or slide (PPTX) or transcript segment (video)."""
    index: int
    text: str


def parse_pdf(file_path: str) -> list[ParsedUnit]:
    units: list[ParsedUnit] = []
    reader = pypdf.PdfReader(file_path)
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            units.append(ParsedUnit(index=i + 1, text=text))
    return units


def parse_pptx(file_path: str) -> list[ParsedUnit]:
    units: list[ParsedUnit] = []
    prs = Presentation(file_path)
    for i, slide in enumerate(prs.slides):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        parts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
        if text:
            units.append(ParsedUnit(index=i + 1, text=text))
    return units


def parse_video_transcript(transcript_text: str, segment_chars: int = 1200) -> list[ParsedUnit]:
    """
    Splits a flat transcript into rough segments. Replace with real
    timestamped ASR segments when a speech-to-text step is wired in.
    """
    units: list[ParsedUnit] = []
    text = transcript_text.strip()
    for i in range(0, len(text), segment_chars):
        segment = text[i : i + segment_chars].strip()
        if segment:
            units.append(ParsedUnit(index=len(units) + 1, text=segment))
    return units


def parse_document(file_path: str, source_type: str) -> list[ParsedUnit]:
    if source_type == "pdf":
        return parse_pdf(file_path)
    if source_type == "pptx":
        return parse_pptx(file_path)
    raise ValueError(f"Unsupported source_type for file parsing: {source_type}")
